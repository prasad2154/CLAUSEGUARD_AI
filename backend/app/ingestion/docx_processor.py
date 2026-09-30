"""
DOCX Processor — Extract text from Word documents.
Preserves heading structure, paragraph order, and section metadata.
"""

import logging
from dataclasses import dataclass, field
from typing import List, Optional

logger = logging.getLogger(__name__)


@dataclass
class DocxParagraph:
    text: str
    style: str          # e.g., 'Heading 1', 'Normal', 'List Bullet'
    is_heading: bool
    level: int          # Heading level 1-9, 0 for body
    page_estimate: int  # Approximate page (word count based)


@dataclass
class DocxExtractionResult:
    paragraphs: List[DocxParagraph]
    full_text: str
    total_paragraphs: int
    estimated_pages: int
    extraction_method: str = "python-docx"
    error: Optional[str] = None


def extract_docx(file_path: str) -> DocxExtractionResult:
    """
    Extract text and structure from a DOCX file.
    Returns paragraphs with heading levels for clause segmentation.
    """
    try:
        from docx import Document
        from docx.shared import Pt

        doc = Document(file_path)
        paragraphs: List[DocxParagraph] = []
        word_count = 0
        words_per_page = 350  # Approximate

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name if para.style else "Normal"
            is_heading = "Heading" in style_name or "Title" in style_name

            # Parse heading level
            level = 0
            if is_heading:
                for i in range(1, 10):
                    if f"Heading {i}" in style_name:
                        level = i
                        break
                if "Title" in style_name:
                    level = 0

            # Estimate page number from word count
            word_count += len(text.split())
            page_estimate = max(1, word_count // words_per_page + 1)

            paragraphs.append(DocxParagraph(
                text=text,
                style=style_name,
                is_heading=is_heading,
                level=level,
                page_estimate=page_estimate,
            ))

        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in table.cells:
                    text = cell.text.strip()
                    if text:
                        word_count += len(text.split())
                        page_estimate = max(1, word_count // words_per_page + 1)
                        paragraphs.append(DocxParagraph(
                            text=text,
                            style="Table Cell",
                            is_heading=False,
                            level=0,
                            page_estimate=page_estimate,
                        ))

        full_text = "\n".join(p.text for p in paragraphs)
        estimated_pages = max(1, word_count // words_per_page + 1)

        return DocxExtractionResult(
            paragraphs=paragraphs,
            full_text=full_text,
            total_paragraphs=len(paragraphs),
            estimated_pages=estimated_pages,
        )

    except Exception as e:
        logger.error(f"DOCX extraction failed: {e}")
        return DocxExtractionResult(
            paragraphs=[],
            full_text="",
            total_paragraphs=0,
            estimated_pages=0,
            error=str(e),
        )
